#!/usr/bin/env python3
"""
Build the parallel Italian-Russian idiom corpus for LeggoManzoni.

The Russian counterpart of build_idiom_parallel.py: same idiom occurrences, same
<note target/targetEnd> join, but against the two Russian editions and writing to
its own artefacts so the English corpus is untouched.

The parsing/matching logic is imported from build_idiom_parallel rather than
reimplemented; only the edition list, the segment column name (ru_segment) and the
output paths differ.

Writes:
    data/idioms_parallel_ru.xlsx     -- the linguists' corpus (1800 rows)
    data/idioms_parallel_ru.json     -- identical records, for scripts
    data/idioms_ru_llm_verify.jsonl  -- minimal input for the alignment-check pass

Makes no LLM calls and never modifies quarantana/ or translations/.

Usage:
    scripts/venv/bin/python scripts/build_idiom_parallel_ru.py
"""

import importlib.util
import json
import logging
import os

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Import the English builder for its shared helpers. Loaded by path because the
# module name is not importable as a package.
_spec = importlib.util.spec_from_file_location(
    "build_idiom_parallel", os.path.join(BASE_DIR, "scripts", "build_idiom_parallel.py")
)
bip = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(bip)

# (edition label, translations/ subdirectory). Chronological; also lexicographic,
# which the row sort relies on.
EDITIONS = [
    ("1936", "Russian_1936"),
    ("1999", "Russian_1999"),
]

COLUMNS = [
    "idiom_id", "label", "tipologia", "parola_chiave",
    "chapter", "comma",
    "it_left", "it_span", "it_right", "it_segment",
    "edition", "translator", "ru_segment",
    "start_id", "end_id", "seg_start_id", "seg_end_id", "note_id", "match_status",
]

# Asserted at build time, mirroring the English builder: a TEI edit that breaks
# alignment must fail loudly rather than quietly shrink the corpus.
EXPECTED_ROWS = 1800
EXPECTED_STATUS = {"ok": 1798, "straddle": 2}

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger(__name__)


def build_rows():
    """One row per idiom occurrence x Russian edition, sorted for reading."""
    with open(os.path.join(BASE_DIR, "data", "concordance.json"), encoding="utf-8") as handle:
        idioms = json.load(handle)["idioms"]

    editions = {}
    for label, dirname in EDITIONS:
        segments, translator, _date = bip.load_edition(dirname)
        editions[label] = (segments, translator)
        log.info("Loaded %s (%s): %d chapters", label, translator, len(segments))

    rows = []
    for idiom_id, idiom in idioms.items():
        for occurrence in idiom["occurrences"]:
            chapter = occurrence["chapter"]
            start = bip.toknum(occurrence["start_id"])
            end = bip.toknum(occurrence["end_id"])

            for label, _dirname in EDITIONS:
                segments, translator = editions[label]
                found = bip.find_segments(segments.get(chapter, []), start, end)
                text, note_id, status = bip.merge_segments(found)

                if found:
                    seg_start = min(s[0] for s in found)
                    seg_end = max(s[1] for s in found)
                    seg_start_id = min(found, key=lambda s: s[0])[4]
                    seg_end_id = max(found, key=lambda s: s[1])[5]
                    it_segment = bip.mark_italian(
                        bip.chapter_tokens(chapter), seg_start, seg_end,
                        occurrence["start_id"], occurrence["end_id"],
                    )
                else:
                    it_segment = seg_start_id = seg_end_id = ""

                rows.append({
                    "idiom_id": int(idiom_id),
                    "label": idiom["label"],
                    "tipologia": idiom["tipologia"],
                    "parola_chiave": idiom["parola_chiave"],
                    "chapter": chapter,
                    "comma": occurrence["comma"],
                    "it_left": " ".join(occurrence["left"]),
                    "it_span": " ".join(occurrence["span_surface"]),
                    "it_right": " ".join(occurrence["right"]),
                    "it_segment": it_segment,
                    "edition": label,
                    "translator": translator,
                    "ru_segment": text,
                    "start_id": occurrence["start_id"],
                    "end_id": occurrence["end_id"],
                    "seg_start_id": seg_start_id,
                    "seg_end_id": seg_end_id,
                    "note_id": note_id,
                    "match_status": status,
                })

    chapter_index = {c: i for i, c in enumerate(bip.CHAPTER_ORDER)}
    rows.sort(key=lambda r: (
        chapter_index[r["chapter"]], bip.toknum(r["start_id"]), r["idiom_id"], r["edition"],
    ))
    return rows


def write_jsonl(rows, path):
    """Minimal input for the alignment-check pass. Makes no LLM calls itself.

    Russian text is written through unchanged (ensure_ascii=False): the ё/е
    spelling of the source must survive verbatim for later exact matching.
    """
    skipped = 0
    with open(path, "w", encoding="utf-8") as handle:
        for row in rows:
            if row["match_status"] == "no_segment":
                skipped += 1
                continue
            record = {
                "row_id": f"{row['idiom_id']}_{row['edition']}",
                "idiom_it": row["label"],
                "it_segment": row["it_segment"],
                "ru_segment": row["ru_segment"],
            }
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")
    log.info("Wrote %d verification lines (%d excluded: no Russian segment)",
             len(rows) - skipped, skipped)


def write_xlsx(rows, path):
    """The linguists' working surface: frozen header, autofilter, wrapped segments."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "idioms"

    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row[column] for column in COLUMNS])

    widths = {
        "label": 28, "parola_chiave": 16, "tipologia": 10,
        "it_left": 30, "it_span": 30, "it_right": 30, "it_segment": 70,
        "translator": 26, "ru_segment": 70,
        "start_id": 14, "end_id": 14, "seg_start_id": 14, "seg_end_id": 14,
        "note_id": 34, "match_status": 14,
    }
    for index, column in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(column, 12)

    wrap = Alignment(wrap_text=True, vertical="top")
    for name in ("it_segment", "ru_segment"):
        column = COLUMNS.index(name) + 1
        for row_cells in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
            row_cells[0].alignment = wrap

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)


def main():
    rows = build_rows()

    counts = {}
    for row in rows:
        counts[row["match_status"]] = counts.get(row["match_status"], 0) + 1
    log.info("Built %d rows: %s", len(rows), counts)

    if len(rows) != EXPECTED_ROWS or counts != EXPECTED_STATUS:
        raise AssertionError(
            f"Alignment changed. Expected {EXPECTED_ROWS} rows {EXPECTED_STATUS}, "
            f"got {len(rows)} rows {counts}. Investigate before regenerating."
        )

    data_dir = os.path.join(BASE_DIR, "data")
    with open(os.path.join(data_dir, "idioms_parallel_ru.json"), "w", encoding="utf-8") as handle:
        json.dump(rows, handle, ensure_ascii=False, indent=2)
    write_jsonl(rows, os.path.join(data_dir, "idioms_ru_llm_verify.jsonl"))
    write_xlsx(rows, os.path.join(data_dir, "idioms_parallel_ru.xlsx"))
    log.info("Wrote 3 artefacts to %s", data_dir)


if __name__ == "__main__":
    main()
