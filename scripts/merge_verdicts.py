#!/usr/bin/env python3
"""
Validate the subagent verdicts and join clean ones onto the corpus.

Validation makes NO LLM calls. It enforces:
  - every en_span is a literal substring of its record's en_segment
  - en_span is non-empty iff found is true
  - confidence is in the allowed set; found is boolean
  - the row_id set equals the input set — no duplicates, no inventions, no missing

Usage:
    scripts/venv/bin/python scripts/merge_verdicts.py
"""

import glob
import json
import logging
import os

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INPUT_PATH = os.path.join(BASE_DIR, "data", "idioms_llm_verify.jsonl")
VERDICT_DIR = os.path.join(BASE_DIR, "data", "verify", "verdicts")
FAILURES_PATH = os.path.join(BASE_DIR, "data", "verify", "failures.jsonl")
CORPUS_PATH = os.path.join(BASE_DIR, "data", "idioms_parallel.json")
XLSX_PATH = os.path.join(BASE_DIR, "data", "idioms_verified.xlsx")

VALID_CONFIDENCE = {"high", "medium", "low"}

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger(__name__)


def validate_verdict(verdict, en_by_row_id):
    """Return a failure reason string, or None if the verdict is well-formed.

    Does not check for duplicates/missing — that is a set-level check in partition().
    """
    row_id = verdict.get("row_id")
    if row_id not in en_by_row_id:
        return "unknown row_id"
    if not isinstance(verdict.get("found"), bool):
        return "found is not boolean"
    if verdict.get("confidence") not in VALID_CONFIDENCE:
        return "bad confidence"
    span = verdict.get("en_span", "")
    if verdict["found"]:
        if not span:
            return "found=true but empty en_span"
        if span not in en_by_row_id[row_id]:
            return "en_span not a literal substring of en_segment"
    else:
        if span:
            return "found=false but non-empty en_span"
    return None


def load_verdicts(directory):
    verdicts = []
    for path in sorted(glob.glob(os.path.join(directory, "verdict_*.jsonl"))):
        with open(path, encoding="utf-8") as handle:
            for line in handle:
                if line.strip():
                    verdicts.append(json.loads(line))
    return verdicts


def partition(verdicts, input_index):
    """Split into (clean, failures). Failures carry a '_failure' reason.

    input_index maps row_id -> en_segment for every expected record.
    """
    clean, failures = [], []
    seen = {}
    for verdict in verdicts:
        reason = validate_verdict(verdict, input_index)
        row_id = verdict.get("row_id")
        if row_id in seen:
            reason = reason or "duplicate row_id"
        seen[row_id] = True
        if reason:
            failures.append({**verdict, "_failure": reason})
        else:
            clean.append(verdict)

    for row_id in input_index:
        if row_id not in seen:
            failures.append({"row_id": row_id, "_failure": "missing verdict"})

    return clean, failures


def write_xlsx(clean, corpus, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    verdict_by_row = {v["row_id"]: v for v in clean}
    columns = [
        "idiom_id", "label", "tipologia", "parola_chiave", "chapter", "comma",
        "it_span", "it_segment", "edition", "translator", "en_segment", "match_status",
        "found", "en_span", "confidence", "reason",
    ]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "verified"
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for row in corpus:
        row_id = f"{row['idiom_id']}_{row['edition']}"
        verdict = verdict_by_row.get(row_id, {})
        sheet.append([
            row["idiom_id"], row["label"], row["tipologia"], row["parola_chiave"],
            row["chapter"], row["comma"], row["it_span"], row["it_segment"],
            row["edition"], row["translator"], row["en_segment"], row["match_status"],
            verdict.get("found", ""), verdict.get("en_span", ""),
            verdict.get("confidence", ""), verdict.get("reason", ""),
        ])

    widths = {"label": 28, "it_segment": 70, "en_segment": 70, "en_span": 40,
              "reason": 40, "translator": 22}
    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(column, 12)
    wrap = Alignment(wrap_text=True, vertical="top")
    for name in ("it_segment", "en_segment", "en_span", "reason"):
        col = columns.index(name) + 1
        for cells in sheet.iter_rows(min_row=2, min_col=col, max_col=col):
            cells[0].alignment = wrap

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)


def main():
    with open(INPUT_PATH, encoding="utf-8") as handle:
        input_index = {
            json.loads(line)["row_id"]: json.loads(line)["en_segment"]
            for line in handle if line.strip()
        }

    verdicts = load_verdicts(VERDICT_DIR)
    clean, failures = partition(verdicts, input_index)

    log.info("Verdicts: %d loaded, %d clean, %d failures", len(verdicts), len(clean), len(failures))

    with open(FAILURES_PATH, "w", encoding="utf-8") as handle:
        for failure in failures:
            handle.write(json.dumps(failure, ensure_ascii=False) + "\n")

    if failures:
        log.warning("%d failures written to %s — re-run those shards before trusting the corpus", len(failures), FAILURES_PATH)
        return

    with open(CORPUS_PATH, encoding="utf-8") as handle:
        corpus = json.load(handle)
    write_xlsx(clean, corpus, XLSX_PATH)
    log.info("Wrote verified corpus to %s", XLSX_PATH)


if __name__ == "__main__":
    main()
